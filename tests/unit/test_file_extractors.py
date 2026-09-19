"""app.rag.file_extractors - text extraction for uploaded knowledge-base
files (spec: help-desk-kit-style KB upload, .md/.txt/.pdf/.docx).
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.rag.file_extractors import FileExtractionError, extract_text


def test_extracts_plain_text():
    assert extract_text("notes.txt", b"hello world") == "hello world"


def test_plain_text_rejects_non_utf8():
    with pytest.raises(FileExtractionError, match="UTF-8"):
        extract_text("notes.txt", b"\xff\xfe not utf-8")


def test_rejects_unsupported_extension():
    with pytest.raises(FileExtractionError, match="Unsupported file extension"):
        extract_text("policy.exe", b"binary junk")


def test_extracts_docx_paragraphs_and_headings():
    document = Document()
    document.add_heading("Return Policy", level=1)
    document.add_paragraph("Customers may return items within 30 days.")
    buf = io.BytesIO()
    document.save(buf)

    text = extract_text("return_policy.docx", buf.getvalue())

    assert "# Return Policy" in text
    assert "Customers may return items within 30 days." in text


def test_extracts_docx_tables_as_gfm_markdown():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Window"
    table.rows[1].cells[0].text = "Electronics"
    table.rows[1].cells[1].text = "14 days"
    buf = io.BytesIO()
    document.save(buf)

    text = extract_text("policy.docx", buf.getvalue())

    # A real GFM table, not just "cell | cell" text - the header separator
    # row is what makes markdown renderers actually draw a table.
    assert "| Item | Window |" in text
    assert "| --- | --- |" in text
    assert "| Electronics | 14 days |" in text


def test_rejects_invalid_docx():
    with pytest.raises(FileExtractionError, match="Could not read this Word document"):
        extract_text("policy.docx", b"not a real docx")


def test_rejects_empty_docx():
    document = Document()
    buf = io.BytesIO()
    document.save(buf)

    with pytest.raises(FileExtractionError, match="No extractable text"):
        extract_text("empty.docx", buf.getvalue())


def test_rejects_invalid_pdf():
    with pytest.raises(FileExtractionError, match="Could not read this PDF"):
        extract_text("policy.pdf", b"not a real pdf")


def test_rejects_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    buf = io.BytesIO()
    writer.write(buf)

    with pytest.raises(FileExtractionError, match="Encrypted"):
        extract_text("policy.pdf", buf.getvalue())


def test_rejects_pdf_with_no_text_layer():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    with pytest.raises(FileExtractionError, match="No extractable text"):
        extract_text("scanned.pdf", buf.getvalue())
